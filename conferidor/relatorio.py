"""Gera o relatório de conferência em Excel, no layout do painel original.

Duas seções (BTG e Bradesco), colunas:
  Fundo | Informado (instituição) | BWAG (CVM) | Diferença (R$) | Diferença (%)
Mais uma aba de detalhe por fundo (dias, período, componentes).
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_ARIAL = "Arial"
_AZUL = "1F4E78"
_CINZA = "D9D9D9"
_VERDE = "C6EFCE"
_VERMELHO = "FFC7CE"
_MOEDA = '#,##0.00;[Red](#,##0.00)'


def _titulo(ws, celula, texto):
    c = ws[celula]
    c.value = texto
    c.font = Font(name=_ARIAL, bold=True, size=12, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=_AZUL)
    c.alignment = Alignment(horizontal="center", vertical="center")


def _cab(ws, linha, cols):
    borda = Border(bottom=Side(style="thin", color="808080"))
    for i, texto in enumerate(cols, start=2):
        c = ws.cell(linha, i, texto)
        c.font = Font(name=_ARIAL, bold=True, size=10)
        c.fill = PatternFill("solid", fgColor=_CINZA)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = borda


def _secao(ws, linha, titulo, resultados, cache):
    """Escreve uma seção (BTG ou Bradesco) a partir da linha dada. Retorna próxima linha livre.

    `cache` acumula {coordenada: valor} das células com fórmula, para injetar
    depois como valor em cache (o arquivo abre já mostrando os números).
    """
    ws.merge_cells(start_row=linha, start_column=2, end_row=linha, end_column=6)
    _titulo(ws, f"B{linha}", titulo)
    linha += 1
    _cab(ws, linha, ["Fundo", "Informado (R$)", "BWAG · CVM (R$)", "Diferença (R$)", "Diferença (%)"])
    linha += 1
    inicio = linha
    soma = {"C": 0.0, "D": 0.0, "E": 0.0}
    for r in resultados:
        ws.cell(linha, 2, r["fundo"]).font = Font(name=_ARIAL, size=10)
        cinf = ws.cell(linha, 3, r["informado"])
        cbwag = ws.cell(linha, 4, r["bwag"])
        cdif = ws.cell(linha, 5, f"=C{linha}-D{linha}" if r["informado"] is not None else None)
        cpct = ws.cell(linha, 6, f"=IFERROR(E{linha}/C{linha},0)" if r["informado"] is not None else None)
        for c in (cinf, cbwag, cdif):
            c.number_format = _MOEDA
            c.font = Font(name=_ARIAL, size=10)
        cpct.number_format = "0.00%"
        cpct.font = Font(name=_ARIAL, size=10)
        if r["diferenca"] is not None:
            # destaca a diferença: verde se ~zero, vermelho se relevante
            cor = _VERDE if abs(r["diferenca"]) <= 1.0 else _VERMELHO
            cdif.fill = PatternFill("solid", fgColor=cor)
            cache[f"E{linha}"] = r["diferenca"]
            cache[f"F{linha}"] = round(r["diferenca"] / r["informado"], 6) if r["informado"] else 0
        for col, val in (("C", r["informado"]), ("D", r["bwag"]), ("E", r["diferenca"])):
            if isinstance(val, (int, float)):
                soma[col] += val
        linha += 1
    # total
    ctot = ws.cell(linha, 2, "TOTAL")
    ctot.font = Font(name=_ARIAL, bold=True, size=10)
    for col in (3, 4, 5):
        L = get_column_letter(col)
        c = ws.cell(linha, col, f"=SUM({L}{inicio}:{L}{linha-1})")
        c.number_format = _MOEDA
        c.font = Font(name=_ARIAL, bold=True, size=10)
        c.fill = PatternFill("solid", fgColor=_CINZA)
        cache[f"{L}{linha}"] = round(soma[L], 2)
    return linha + 2


def gerar(resultados, ano, mes, caminho):
    """Gera o arquivo Excel de conferência em `caminho`."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Conferência"
    ws.sheet_view.showGridLines = False

    meses = ["", "Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez"]
    ws.merge_cells("B1:F1")
    _titulo(ws, "B1", f"Conferência de Receita — competência {meses[mes]}/{ano}")
    ws.row_dimensions[1].height = 22

    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 34
    for col in "CDEF":
        ws.column_dimensions[col].width = 17

    btg = [r for r in resultados if r["instituicao"] == "BTG"]
    brad = [r for r in resultados if r["instituicao"] == "BRADESCO"]

    cache = {}
    linha = 3
    if btg:
        linha = _secao(ws, linha, "GESTÃO — BTG", btg, cache)
    if brad:
        linha = _secao(ws, linha, "GESTÃO — BRADESCO", brad, cache)

    nota = ws.cell(
        linha + 1, 2,
        "Coluna 'Informado' = valor enviado pela instituição (preencher/colar). "
        "Coluna 'BWAG · CVM' = recalculado automaticamente a partir do Informe Diário da CVM.",
    )
    nota.font = Font(name=_ARIAL, italic=True, size=8, color="808080")

    _aba_detalhe(wb, resultados)
    cache_mem = _aba_memoria(wb, resultados)

    import os
    os.makedirs(os.path.dirname(os.path.abspath(caminho)), exist_ok=True)
    wb.save(caminho)
    # Grava os valores em cache das fórmulas (o arquivo abre mostrando os números,
    # mesmo antes de o Excel recalcular). As fórmulas continuam vivas.
    _injetar_cache(caminho, {"Conferência": cache, "Memória de Cálculo": cache_mem})
    return caminho


def _injetar_cache(caminho, valores_por_aba):
    """Injeta valores em cache (<v>) nas células com fórmula, via patch do XML.

    openpyxl não grava o valor calculado junto da fórmula; este passo o adiciona,
    para o arquivo abrir já exibindo os números (as fórmulas permanecem).
    """
    import re
    import zipfile

    wb = load_workbook_names(caminho)
    with zipfile.ZipFile(caminho) as zin:
        conteudo = {n: zin.read(n) for n in zin.namelist()}

    for idx, titulo in enumerate(wb, start=1):
        cache = valores_por_aba.get(titulo)
        if not cache:
            continue
        chave = f"xl/worksheets/sheet{idx}.xml"
        if chave not in conteudo:
            continue
        xml = conteudo[chave].decode("utf-8")
        for coord, val in cache.items():
            padrao = re.compile(r'(<c r="%s"[^>]*>)(.*?)(</c>)' % re.escape(coord), re.DOTALL)

            def _sub(m):
                interno = re.sub(r"<v\s*/>|<v>.*?</v>", "", m.group(2), flags=re.DOTALL)
                return f"{m.group(1)}{interno}<v>{val}</v>{m.group(3)}"

            xml = padrao.sub(_sub, xml, count=1)
        conteudo[chave] = xml.encode("utf-8")

    import os
    import time

    tmp = caminho + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for nome, dados in conteudo.items():
            zout.writestr(nome, dados)
    # No Windows o destino pode estar travado por um instante (Excel aberto ou o
    # Google Drive sincronizando). Tenta algumas vezes; se não der, mantém o
    # arquivo já salvo (o Excel recalcula as fórmulas ao abrir) e não quebra.
    for tentativa in range(6):
        try:
            os.replace(tmp, caminho)
            return
        except PermissionError:
            if tentativa < 5:
                time.sleep(0.5)
    try:
        os.remove(tmp)
    except OSError:
        pass


def load_workbook_names(caminho):
    """Retorna os nomes das abas na ordem (sheet1.xml, sheet2.xml, ...)."""
    from openpyxl import load_workbook

    return load_workbook(caminho, read_only=True).sheetnames


def _aba_detalhe(wb, resultados):
    ws = wb.create_sheet("Detalhe")
    ws.sheet_view.showGridLines = False
    cols = ["Fundo", "CNPJ", "Instituição", "Regra", "Dias úteis",
            "Início período", "Fim período", "Gestão (R$)", "Controladoria (R$)",
            "Cogestão (R$)", "Extra (R$)", "Líquido BWAG (R$)"]
    for i, t in enumerate(cols, start=1):
        c = ws.cell(1, i, t)
        c.font = Font(name=_ARIAL, bold=True, size=10, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=_AZUL)
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    for j, r in enumerate(resultados, start=2):
        comp = r["componentes"]
        vals = [
            r["fundo"], r["cnpj"], r["instituicao"], r["regra"], r["dias"],
            r["data_inicio"], r["data_fim"],
            comp.get("gestao"), comp.get("controladoria"),
            comp.get("cogestao"), comp.get("extra"), r["bwag"],
        ]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(j, i, v)
            c.font = Font(name=_ARIAL, size=10)
            if i >= 8:
                c.number_format = _MOEDA
    widths = [30, 20, 12, 26, 10, 14, 14, 14, 16, 14, 12, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _br_data(iso):
    """'AAAA-MM-DD' -> 'DD/MM/AAAA' (vazio se None)."""
    if not iso:
        return ""
    a, m, d = iso.split("-")
    return f"{d}/{m}/{a}"


_COTA_FMT = "#,##0.00000000"


def _aba_memoria(wb, resultados):
    """Memória de cálculo dia a dia: por fundo, uma linha por dia (PL, cota, ganho),
    o subtotal de gestão e o fechamento até a receita líquida BWAG.

    Feita para ser lida por qualquer pessoa e para bater com a base de cálculo do banco.
    """
    ws = wb.create_sheet("Memória de Cálculo")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 22

    borda = Border(bottom=Side(style="thin", color="808080"))
    linha = 2
    cache_mem = {}  # {coordenada: valor} dos subtotais (fórmulas), para o cache
    for r in resultados:
        memoria = r.get("memoria") or []

        # Título do fundo
        ws.merge_cells(start_row=linha, start_column=2, end_row=linha, end_column=6)
        _titulo(ws, f"B{linha}", r["fundo"])
        ws.row_dimensions[linha].height = 20
        linha += 1

        # Linha de contexto (CNPJ · instituição · regra · período)
        periodo = f"{_br_data(r.get('data_inicio'))} a {_br_data(r.get('data_fim'))}"
        ctx = (
            f"CNPJ {r['cnpj']}  ·  {r['instituicao']}  ·  regra: {r['regra']}  ·  "
            f"{r['dias']} dias úteis  ·  período {periodo}"
        )
        c = ws.cell(linha, 2, ctx)
        ws.merge_cells(start_row=linha, start_column=2, end_row=linha, end_column=6)
        c.font = Font(name=_ARIAL, italic=True, size=8, color="595959")
        linha += 1

        if not memoria:
            c = ws.cell(linha, 2, "Sem dados da CVM para o período.")
            c.font = Font(name=_ARIAL, italic=True, size=9, color="C00000")
            linha += 2
            continue

        # Cabeçalho das colunas do dia a dia
        for i, texto in enumerate(
            ["Data", "PL do dia (R$)", "Cota", "Taxa gestão", "Ganho de gestão no dia (R$)"],
            start=2
        ):
            hc = ws.cell(linha, i, texto)
            hc.font = Font(name=_ARIAL, bold=True, size=10)
            hc.fill = PatternFill("solid", fgColor=_CINZA)
            hc.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            hc.border = borda
        linha += 1

        inicio = linha
        for d in memoria:
            ws.cell(linha, 2, _br_data(d["data"])).font = Font(name=_ARIAL, size=10)
            cpl = ws.cell(linha, 3, d["pl"])
            cpl.number_format = _MOEDA
            cpl.font = Font(name=_ARIAL, size=10)
            cct = ws.cell(linha, 4, d.get("cota"))
            cct.number_format = _COTA_FMT
            cct.font = Font(name=_ARIAL, size=10)
            ctx = ws.cell(linha, 5, d.get("taxa"))
            ctx.number_format = "0.00%"
            ctx.font = Font(name=_ARIAL, size=10)
            ctx.alignment = Alignment(horizontal="center")
            cg = ws.cell(linha, 6, d["ganho_gestao"])
            cg.number_format = _MOEDA
            cg.font = Font(name=_ARIAL, size=10)
            linha += 1

        # Subtotal de gestão (soma dos ganhos diários)
        cs = ws.cell(linha, 2, "Gestão (soma dos dias)")
        cs.font = Font(name=_ARIAL, bold=True, size=10)
        cg = ws.cell(linha, 6, f"=SUM(F{inicio}:F{linha-1})")
        cg.number_format = _MOEDA
        cg.font = Font(name=_ARIAL, bold=True, size=10)
        cg.fill = PatternFill("solid", fgColor=_CINZA)
        cache_mem[f"F{linha}"] = round(sum(d["ganho_gestao"] for d in memoria), 2)
        linha += 1

        # Fechamento: componentes que reduzem a receita e o líquido final
        comp = r.get("componentes") or {}
        fechamento = []
        if comp.get("cogestao"):
            fechamento.append(("(−) Cogestão (piso mensal)", -comp["cogestao"]))
        if comp.get("extra"):
            fechamento.append(("(−) Taxa extra (piso mensal)", -comp["extra"]))
        if comp.get("controladoria"):
            fechamento.append(("(−) Controladoria", -comp["controladoria"]))
        for rotulo, valor in fechamento:
            cl = ws.cell(linha, 2, rotulo)
            cl.font = Font(name=_ARIAL, size=10, color="595959")
            cv = ws.cell(linha, 6, valor)
            cv.number_format = _MOEDA
            cv.font = Font(name=_ARIAL, size=10, color="595959")
            linha += 1

        # (não começar o texto com "=": o openpyxl trataria como fórmula -> #NAME? no Excel)
        cl = ws.cell(linha, 2, "Receita líquida BWAG")
        cl.font = Font(name=_ARIAL, bold=True, size=10, color="FFFFFF")
        cl.fill = PatternFill("solid", fgColor=_AZUL)
        cv = ws.cell(linha, 6, r["bwag"])
        cv.number_format = _MOEDA
        cv.font = Font(name=_ARIAL, bold=True, size=10, color="FFFFFF")
        cv.fill = PatternFill("solid", fgColor=_AZUL)
        linha += 2  # espaço entre fundos
    return cache_mem
